import io
import time
import zipfile
from dataclasses import asdict, dataclass, field
from typing import Any

from .statuses import ExtractionStatus

MAX_EXTRACTION_TIME_SECONDS = 60
MAX_PAGES = 200
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 MiB
MAX_EXTRACTED_CHARS = 5 * 1024 * 1024  # 5 MB


@dataclass
class ExtractionResult:
    status: str
    extracted_text: str = ""
    page_count: int = 0
    metadata: dict[str, Any] = field(default_factory=dict)
    error_message: str = ""

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _count_words(text: str) -> int:
    return len(text.split()) if text else 0


def extract_pdf(file_bytes: bytes) -> ExtractionResult:
    """
    Trích xuất văn bản từ tệp PDF bằng pypdf với giới hạn tài nguyên
    và phân rã provenance theo trang.
    """
    try:
        import pypdf
        from pypdf.errors import PdfReadError
    except ImportError:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message="Thư viện pypdf chưa được cài đặt trong môi trường worker.",
        )

    start_time = time.perf_counter()
    try:
        stream = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(stream, strict=False)

        if reader.is_encrypted:
            try:
                # Thử giải mã với mật khẩu rỗng (một số PDF mã hóa rỗng tiêu chuẩn)
                decrypted = reader.decrypt("")
                if decrypted == pypdf.PasswordType.NOT_DECRYPTED:
                    return ExtractionResult(
                        status=ExtractionStatus.FAILED,
                        error_message="Tệp PDF có mật khẩu bảo vệ hoặc được mã hóa không thể đọc.",
                        metadata={"encrypted": True},
                    )
            except Exception:
                return ExtractionResult(
                    status=ExtractionStatus.FAILED,
                    error_message="Tệp PDF có mật khẩu bảo vệ hoặc mã hóa không hợp lệ.",
                    metadata={"encrypted": True},
                )

        total_pages = len(reader.pages)
        if total_pages == 0:
            return ExtractionResult(
                status=ExtractionStatus.NO_TEXT,
                page_count=0,
                error_message="Tệp PDF không chứa trang nào.",
                metadata={"total_pages": 0, "parser": "pypdf"},
            )

        truncated = False
        pages_to_process = total_pages
        if total_pages > MAX_PAGES:
            pages_to_process = MAX_PAGES
            truncated = True

        pages_data: list[dict[str, Any]] = []
        full_text_parts: list[str] = []
        total_chars = 0

        for idx in range(pages_to_process):
            page = reader.pages[idx]
            page_text = page.extract_text() or ""
            page_text = page_text.strip()

            p_chars = len(page_text)
            p_words = _count_words(page_text)

            pages_data.append(
                {
                    "page_number": idx + 1,
                    "char_count": p_chars,
                    "word_count": p_words,
                    "text": page_text[:1000],  # Lưu mẫu snippet cho provenance
                }
            )

            if page_text:
                full_text_parts.append(f"--- [Trang {idx + 1}] ---\n{page_text}")
                total_chars += p_chars

            if total_chars >= MAX_EXTRACTED_CHARS:
                truncated = True
                break

        full_text = "\n\n".join(full_text_parts).strip()
        duration_ms = round((time.perf_counter() - start_time) * 1000, 2)

        meta = {
            "parser": "pypdf",
            "total_pages": total_pages,
            "processed_pages": len(pages_data),
            "total_chars": len(full_text),
            "total_words": _count_words(full_text),
            "truncated": truncated,
            "duration_ms": duration_ms,
            "pages": pages_data,
        }

        if not full_text:
            return ExtractionResult(
                status=ExtractionStatus.NO_TEXT,
                page_count=total_pages,
                metadata=meta,
                error_message=(
                    "Không tìm thấy luồng ký tự văn bản có thể trích xuất "
                    "(có thể là tệp scan ảnh thuần; OCR hiện chưa được bật)."
                ),
            )

        return ExtractionResult(
            status=ExtractionStatus.SUCCESS,
            extracted_text=full_text[:MAX_EXTRACTED_CHARS],
            page_count=total_pages,
            metadata=meta,
        )

    except PdfReadError as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=f"Tệp PDF bị hỏng hoặc cấu trúc không hợp lệ: {exc}",
            metadata={"parser": "pypdf", "corrupted": True},
        )
    except Exception as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=f"Lỗi không xác định khi trích xuất PDF: {exc}",
            metadata={"parser": "pypdf"},
        )


def extract_docx(file_bytes: bytes) -> ExtractionResult:
    """
    Trích xuất văn bản từ tệp DOCX với cấu trúc đoạn văn, tiêu đề và bảng.
    """
    try:
        import docx
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message="Thư viện python-docx chưa được cài đặt trong môi trường worker.",
        )

    start_time = time.perf_counter()
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(item.file_size for item in entries) > 100 * 1024 * 1024
                or any(item.flag_bits & 1 for item in entries)
            ):
                return ExtractionResult(
                    status=ExtractionStatus.FAILED, error_message="DOCX archive exceeds safe limits"
                )
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        sections_data: list[dict[str, Any]] = []
        full_text_parts: list[str] = []
        sec_idx = 1
        truncated = False

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else "Normal"
            sections_data.append(
                {
                    "section_index": sec_idx,
                    "style": style_name,
                    "char_count": len(text),
                    "word_count": _count_words(text),
                    "snippet": text[:200],
                }
            )
            full_text_parts.append(text)
            sec_idx += 1

            if len("\n".join(full_text_parts)) >= MAX_EXTRACTED_CHARS:
                truncated = True
                break

        # Trích xuất bảng nếu chưa vượt trần
        if not truncated:
            for t_idx, table in enumerate(doc.tables, 1):
                table_rows_text: list[str] = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_rows_text.append(" | ".join(row_cells))
                if table_rows_text:
                    t_str = f"--- [Bảng {t_idx}] ---\n" + "\n".join(table_rows_text)
                    full_text_parts.append(t_str)

        full_text = "\n\n".join(full_text_parts).strip()
        duration_ms = round((time.perf_counter() - start_time) * 1000, 2)

        meta = {
            "parser": "python-docx",
            "section_count": len(sections_data),
            "table_count": len(doc.tables),
            "total_chars": len(full_text),
            "total_words": _count_words(full_text),
            "truncated": truncated,
            "duration_ms": duration_ms,
            "sections": sections_data[:100],  # Lưu 100 section đầu để tránh bloat JSON
        }

        if not full_text:
            return ExtractionResult(
                status=ExtractionStatus.NO_TEXT,
                page_count=1,
                metadata=meta,
                error_message="Tệp DOCX rỗng hoặc không chứa nội dung văn bản.",
            )

        return ExtractionResult(
            status=ExtractionStatus.SUCCESS,
            extracted_text=full_text[:MAX_EXTRACTED_CHARS],
            page_count=0,
            metadata=meta,
        )

    except (PackageNotFoundError, zipfile.BadZipFile) as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=f"Tệp DOCX bị lỗi hoặc không phải định dạng nén hợp lệ: {exc}",
            metadata={"parser": "python-docx", "corrupted": True},
        )
    except Exception as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=f"Lỗi khi trích xuất DOCX: {exc}",
            metadata={"parser": "python-docx"},
        )


def extract_txt(file_bytes: bytes) -> ExtractionResult:
    """
    Trích xuất văn bản từ tệp văn bản thuần (TXT, MD, CSV) với giải mã an toàn nhiều bảng mã.
    """
    start_time = time.perf_counter()
    decoded_text = ""
    used_encoding = ""

    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            decoded_text = file_bytes.decode(enc)
            used_encoding = enc
            break
        except UnicodeDecodeError:
            continue

    if not used_encoding:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=(
                "Không thể giải mã văn bản với các bảng mã thông dụng "
                "(UTF-8, UTF-16, CP1252, Latin-1)."
            ),
            metadata={"parser": "plain_text"},
        )

    clean_text = decoded_text.strip()
    duration_ms = round((time.perf_counter() - start_time) * 1000, 2)

    meta = {
        "parser": "plain_text",
        "encoding": used_encoding,
        "total_chars": len(clean_text),
        "total_words": _count_words(clean_text),
        "duration_ms": duration_ms,
    }

    if not clean_text:
        return ExtractionResult(
            status=ExtractionStatus.NO_TEXT,
            page_count=1,
            metadata=meta,
            error_message="Tệp văn bản rỗng.",
        )

    return ExtractionResult(
        status=ExtractionStatus.SUCCESS,
        extracted_text=clean_text[:MAX_EXTRACTED_CHARS],
        page_count=0,
        metadata=meta,
    )


def extract_file_content(
    file_bytes: bytes,
    file_name: str,
    mime_type: str = "",
) -> ExtractionResult:
    """
    Điểm điều phối trích xuất văn bản theo phần mở rộng và mime type.
    Thực thi kiểm tra kích thước trần 50MB và giới hạn tài nguyên.
    """
    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            error_message=(
                f"Kích thước tệp ({len(file_bytes)} bytes) vượt quá "
                "giới hạn trích xuất tối đa (50MB)."
            ),
            metadata={"file_size": len(file_bytes), "max_size": MAX_FILE_SIZE_BYTES},
        )

    if not file_bytes:
        return ExtractionResult(
            status=ExtractionStatus.NO_TEXT,
            page_count=0,
            error_message="Tệp rỗng (0 bytes).",
            metadata={"file_size": 0},
        )

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    mime = (mime_type or "").lower()

    if ext == "pdf" or "application/pdf" in mime:
        return extract_pdf(file_bytes)
    elif (
        ext in ("docx", "dotx")
        or "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in mime
    ):
        return extract_docx(file_bytes)
    elif ext in ("txt", "md", "csv", "log") or "text/plain" in mime or "text/markdown" in mime:
        return extract_txt(file_bytes)
    else:
        return ExtractionResult(
            status=ExtractionStatus.UNSUPPORTED,
            error_message=(
                f"Định dạng tệp '{ext or 'unknown'}' ({mime or 'unknown mime'}) "
                "chưa được hỗ trợ trích xuất văn bản."
            ),
            metadata={"extension": ext, "mime_type": mime},
        )
